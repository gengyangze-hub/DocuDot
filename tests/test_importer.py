"""导入解析器测试：四种写法、容错边界、真实样例文件。"""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest

from app.errors import ValidationFailed
from app.services.importer import (
    detect_format,
    parse_delimited,
    parse_excel,
    parse_inline_entry,
    parse_input,
    parse_json,
    parse_markdown,
)

SAMPLES = Path(__file__).resolve().parent.parent / "samples"

#: component-inventory（模糊匹配的参考项目）的导出结构
COMPONENT_INVENTORY_JSON = json.dumps(
    {
        "schema_version": 1,
        "app": "component-inventory",
        "created_at": "2026-10-03T23:46:55+08:00",
        "next_seq": 3,
        "components": [
            {
                "id": "1fc4e867",
                "seq": 1,
                "tags": ["C", "1uF", "16V", "MLCC", "0805"],
                "stock": {"mode": "coarse", "level": 3},
                "note": "",
            },
            {
                "id": "a5297305",
                "seq": 2,
                "tags": ["R", "10kΩ", "0805"],
                "stock": {"mode": "coarse", "level": 2},
                "note": "",
            },
        ],
    },
    ensure_ascii=False,
)


# --------------------------------------------------------------------------- #
# Markdown 表格
# --------------------------------------------------------------------------- #
TABLE_MD = """\
# 库存清单

## STM元器件

| 名称 | 数量 | 位置 | 规格 | 别名 | 备注 |
| --- | --- | --- | --- | --- | --- |
| STM32F103C8T6 | 25 | A柜-1层-盒3 | LQFP48 | F103C8, STM32F103 | 蓝药丸 |
| 0.1uF 50V MLCC | 500 | A柜-2层-盒1 | 0805 | 104, 100nF | |

## Steam游戏卡

| 名称 | 数量 | 位置 | 规格 | 别名 | 备注 |
| --- | --- | --- | --- | --- | --- |
| Steam 50元充值卡 | 10 | B柜-抽屉1 | 50元 | 50元卡 | |
"""


def test_markdown_table() -> None:
    preview = parse_markdown(TABLE_MD)
    assert preview.total == 3
    first = preview.rows[0]
    assert first.name == "STM32F103C8T6"
    assert first.quantity == 25
    assert first.location == "A柜-1层-盒3"
    assert first.spec == "LQFP48"
    assert first.aliases == ["F103C8", "STM32F103"]
    assert first.category == "STM元器件"      # Markdown 标题直接当分类名
    assert preview.rows[2].category == "Steam游戏卡"


def test_table_header_wording_variants() -> None:
    """表头用不同措辞也要认得出来，列顺序可以打乱。"""
    text = """\
| 品名 | 存放位置 | 库存 | 封装 | 标签 |
| --- | --- | --- | --- | --- |
| NE555 | C柜-1层 | 30 | DIP-8 | 555, 定时器 |
"""
    preview = parse_markdown(text)
    assert preview.total == 1
    row = preview.rows[0]
    assert (row.name, row.quantity, row.location, row.spec) == ("NE555", 30, "C柜-1层", "DIP-8")
    assert row.aliases == ["555", "定时器"]
    assert set(preview.detected_columns) >= {"品名", "存放位置", "库存", "封装", "标签"}


def test_table_without_header_falls_back_to_positional() -> None:
    text = "| STM32F103C8T6 | 25 | A柜 | STM元器件 | LQFP48 | F103C8 | 备注 |\n"
    preview = parse_markdown(text)
    assert preview.total == 1
    assert preview.rows[0].quantity == 25
    assert any("固定列序" in w for w in preview.warnings)


# --------------------------------------------------------------------------- #
# 行内式 / 键值块
# --------------------------------------------------------------------------- #
def test_inline_entries() -> None:
    text = """\
STM32F103C8T6 × 25 @A柜-1层-盒3 #F103C8 #STM32F103 (LQFP48)
NE555 30个 @C柜-1层 #555 (DIP-8)
"""
    preview = parse_markdown(text)
    assert preview.total == 2
    stm = preview.rows[0]
    assert stm.name == "STM32F103C8T6"
    assert stm.quantity == 25
    assert stm.location == "A柜-1层-盒3"
    assert stm.aliases == ["F103C8", "STM32F103"]
    assert stm.spec == "LQFP48"
    assert preview.rows[1].quantity == 30


def test_inline_trailing_quantity() -> None:
    """行尾裸数量：名称里带小数点也不能截错。"""
    from app.services.importer import parse_inline_entry

    row = parse_inline_entry("AMS1117-3.3 5")
    assert row is not None
    assert row.name == "AMS1117-3.3"
    assert row.quantity == 5

    row = parse_inline_entry("0.1uF 50V MLCC（DIP-8） 100")
    assert row is not None
    assert row.name == "0.1uF 50V MLCC"      # 规格被正确剥掉
    assert row.spec == "DIP-8"
    assert row.quantity == 100

    # 前导零不是数量（封装码）
    row = parse_inline_entry("0805")
    assert row is not None
    assert row.quantity == 0


def test_numbered_list_marker_is_stripped() -> None:
    """实测 bug：采购清单带序号时，序号被当成名称的一部分，结果全部匹配不上。"""
    for line, name in [
        ("1. ESP32-C3 × 100", "ESP32-C3"),
        ("12. 4.7kΩ（0805） × 200", "4.7kΩ"),
        ("3) 排母 11P × 200", "排母 11P"),
        ("4、1kΩ × 200", "1kΩ"),
        ("（5）10kΩ × 200", "10kΩ"),
        ("① 51Ω × 200", "51Ω"),
        ("- 220Ω × 200", "220Ω"),
    ]:
        row = parse_inline_entry(line)
        assert row is not None, line
        assert row.name == name, f"{line} → {row.name!r}"


def test_number_marker_does_not_eat_real_names() -> None:
    """序号剥离不能误伤「4.7kΩ」这种真名称 —— 所以要求序号后必须有空白。"""
    for line, name in [
        ("4.7kΩ（0805） × 200", "4.7kΩ"),
        ("0.1uF 50V × 100", "0.1uF 50V"),
        ("2N3904 × 50", "2N3904"),
        ("10kΩ 0805 200", "10kΩ"),
    ]:
        row = parse_inline_entry(line)
        assert row is not None, line
        assert row.name == name, f"{line} → {row.name!r}"


def test_dimension_spec_does_not_steal_quantity() -> None:
    """实测 bug：``5x11mm`` 里的 ``x11`` 被当成乘号，数量解析成 11。"""
    row = parse_inline_entry("10uF 25V（铝电解 5x11mm） × 200")
    assert row is not None
    assert row.quantity == 200          # 不是 11
    assert row.spec == "铝电解 5x11mm"
    assert row.name == "10uF 25V"

    # 正常写法仍然有效
    assert parse_inline_entry("100nF × 500").quantity == 500
    assert parse_inline_entry("100nF x500").quantity == 500


def test_key_value_blocks() -> None:
    text = """\
### 1N4148
- 数量: 200
- 位置: C柜-2层
- 规格: DO-35
- 别名: 4148

### Steam 50元充值卡
- 数量: 10
- 位置: B柜-抽屉1
"""
    preview = parse_markdown(text)
    assert preview.total == 2
    assert preview.rows[0].name == "1N4148"
    assert preview.rows[0].quantity == 200
    # ``###`` 是条目名不是分类 —— 标题不能被当成分类名
    assert preview.rows[1].category == "未分类"
    assert preview.rows[1].name == "Steam 50元充值卡"


# --------------------------------------------------------------------------- #
# 回归：说明性文字不能变成物品
# --------------------------------------------------------------------------- #
def test_blockquote_and_prose_are_skipped() -> None:
    """引用块、分隔线、说明句子都不应产生条目（曾经被误当成物品导入）。"""
    text = """\
# 库存清单

> 这份文件本身就是可导入格式的范例：上传即可。

---

下面是一些随手写的说明文字，没有数量也没有位置。

| 名称 | 数量 | 位置 |
| --- | --- | --- |
| STM32F103C8T6 | 25 | A柜 |
"""
    preview = parse_markdown(text)
    assert preview.total == 1
    assert preview.rows[0].name == "STM32F103C8T6"


def test_prose_only_markdown_yields_nothing() -> None:
    text = """\
# 说明

这是一份纯说明文档，只有文字，没有任何库存数据。
换行也还是说明。
"""
    preview = parse_markdown(text)
    assert preview.total == 0


# --------------------------------------------------------------------------- #
# 真实样例文件
# --------------------------------------------------------------------------- #
def test_sample_inventory_md_is_clean() -> None:
    """samples/sample_inventory.md 必须恰好导入 10 条，不多不少。"""
    content = (SAMPLES / "sample_inventory.md").read_text(encoding="utf-8")
    preview = parse_input("sample_inventory.md", text=content)
    names = [row.name for row in preview.rows]
    assert preview.total == 10, names
    assert "STM32F103C8T6" in names
    assert not any(name.startswith(">") for name in names)
    assert preview.warnings == []


def test_sample_components_txt() -> None:
    content = (SAMPLES / "sample_components.txt").read_text(encoding="utf-8")
    preview = parse_input("sample_components.txt", text=content)
    assert preview.total == 7, [row.name for row in preview.rows]
    assert preview.rows[0].name == "STM32F103C8T6"
    assert preview.rows[0].quantity == 25


# --------------------------------------------------------------------------- #
# CSV / Excel
# --------------------------------------------------------------------------- #
def test_csv_with_header() -> None:
    text = "名称,数量,位置,封装\nNE555,30,C柜-1层,DIP-8\n"
    preview = parse_delimited(text, delimiter=",")
    assert preview.total == 1
    assert preview.rows[0].spec == "DIP-8"


def test_excel_multiple_sheets() -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "元件"
    sheet.append(["名称", "数量", "位置", "封装", "别名"])
    sheet.append(["NE555", 30, "C柜-1层", "DIP-8", "555"])

    sheet2 = workbook.create_sheet("游戏卡")
    sheet2.append(["品名", "库存", "存放位置", "标签"])
    sheet2.append(["Steam 200元充值卡", 3, "B柜-抽屉3", "200元卡"])

    buffer = io.BytesIO()
    workbook.save(buffer)

    preview = parse_excel(buffer.getvalue())
    assert preview.total == 2
    categories = {row.name: row.category for row in preview.rows}
    assert categories["NE555"] == "元件"                 # 工作表名直接当分类名
    assert categories["Steam 200元充值卡"] == "游戏卡"


def test_excel_without_header_row() -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["NE555", 30, "C柜-1层"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    preview = parse_excel(buffer.getvalue())
    assert preview.total == 1
    assert preview.rows[0].name == "NE555"
    assert preview.rows[0].quantity == 30
    assert any("没识别到表头" in w for w in preview.warnings)


# --------------------------------------------------------------------------- #
# 格式判定与编码
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    ("filename", "expected"),
    [
        ("a.xlsx", "excel"),
        ("a.xlsm", "excel"),
        ("a.csv", "csv"),
        ("a.md", "markdown"),
        ("a.markdown", "markdown"),
        ("a.txt", "text"),
        ("a.json", "json"),
    ],
)
def test_detect_format(filename: str, expected: str) -> None:
    assert detect_format(filename) == expected


def test_xls_is_rejected_with_helpful_message() -> None:
    from app.errors import ValidationFailed

    with pytest.raises(ValidationFailed) as exc:
        detect_format("old.xls")
    assert "另存为" in exc.value.message


def test_gbk_encoded_file() -> None:
    data = "名称,数量,位置\nNE555,30,C柜\n".encode("gbk")
    preview = parse_input("a.csv", data=data)
    assert preview.total == 1
    assert preview.rows[0].name == "NE555"


# --------------------------------------------------------------------------- #
# JSON
# --------------------------------------------------------------------------- #
def test_parse_component_inventory_json() -> None:
    """component-inventory 的导出：身份是 tags，库存是粗粒度等级。"""
    preview = parse_input("inventory.json", text=COMPONENT_INVENTORY_JSON)
    assert preview.fmt == "json"
    assert preview.total == 2

    first = preview.rows[0]
    # 封装单独进 spec，名称里不重复写
    assert first.name == "C 1uF 16V"
    assert first.spec == "MLCC 0805"
    assert first.quantity == 200              # level 3 → 200
    assert first.category == "未分类"          # 这个格式不带分类，交给 AI
    assert "原粗粒度库存 level 3" in first.note
    assert preview.warnings                    # 有映射说明

    second = preview.rows[1]
    assert second.name == "R 10kΩ"
    assert second.spec == "0805"
    assert second.quantity == 50              # level 2 → 50


def test_parse_component_inventory_exact_stock() -> None:
    text = json.dumps(
        {"components": [{"tags": ["R", "1k", "0805"], "stock": {"mode": "exact", "quantity": 137}}]},
        ensure_ascii=False,
    )
    preview = parse_json(text)
    assert preview.rows[0].quantity == 137    # 精确数量不做映射


def test_parse_generic_json_array() -> None:
    text = json.dumps(
        [
            {"名称": "NE555", "数量": 30, "位置": "C柜-1层", "封装": "DIP-8", "别名": ["555", "定时器"]},
            {"name": "AMS1117", "qty": 50, "category": "STM元器件"},
        ],
        ensure_ascii=False,
    )
    preview = parse_json(text)
    assert preview.total == 2
    first = preview.rows[0]
    assert first.name == "NE555"
    assert (first.quantity, first.location, first.spec) == (30, "C柜-1层", "DIP-8")
    assert first.aliases == ["555", "定时器"]
    assert preview.rows[1].quantity == 50


def test_parse_json_with_items_wrapper() -> None:
    text = json.dumps({"total": 1, "items": [{"name": "杜邦线", "quantity": 100}]}, ensure_ascii=False)
    preview = parse_json(text)
    assert preview.total == 1
    assert preview.rows[0].quantity == 100


def test_parse_json_plain_strings() -> None:
    preview = parse_json('["NE555", "AMS1117-3.3"]')
    assert [row.name for row in preview.rows] == ["NE555", "AMS1117-3.3"]


def test_parse_json_skips_nameless_entries() -> None:
    text = json.dumps([{"quantity": 5}, {"name": "NE555", "quantity": 1}], ensure_ascii=False)
    preview = parse_json(text)
    assert preview.total == 1
    assert preview.warnings


@pytest.mark.parametrize("text", ["{not json", '{"foo": 1}', "[]", '{"items": []}'])
def test_parse_json_rejects_unusable_input(text: str) -> None:
    with pytest.raises(ValidationFailed):
        parse_json(text)


def test_json_detected_by_content_without_extension() -> None:
    assert detect_format("upload", text='{"items": []}') == "json"
    assert detect_format("upload", text="   [{...}]") == "json"
    # 普通文本不能被误判
    assert detect_format("upload", text="NE555 30个 @C柜") == "text"


# --------------------------------------------------------------------------- #
# 裸写的封装码要提到 spec（实测 bug）
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    ("text", "name", "spec", "quantity"),
    [
        ("100Ω电容 0805 100", "100Ω电容", "0805", 100),          # 实测原文
        ("1uF 16V MLCC 0805", "1uF 16V MLCC", "0805", 0),
        ("NE555 DIP-8 30", "NE555", "DIP-8", 30),
        ("AMS1117-3.3 SOT-223 50", "AMS1117-3.3", "SOT-223", 50),
        ("排针 直插 100", "排针", "直插", 100),
        ("NE555（DIP-8） 30", "NE555", "DIP-8", 30),              # 括号写法优先
        # 不该误伤
        ("USB线 2.0 3", "USB线 2.0", "", 3),
        ("0805", "0805", "", 0),
        ("杜邦线 100", "杜邦线", "", 100),
    ],
)
def test_bare_package_is_promoted_to_spec(text: str, name: str, spec: str, quantity: float) -> None:
    """名称末尾裸写的封装码要成为规格，否则会重复追问一次封装。"""
    from app.services.importer import parse_inline_entry

    row = parse_inline_entry(text)
    assert row is not None, text
    assert (row.name, row.spec, row.quantity) == (name, spec, quantity)
