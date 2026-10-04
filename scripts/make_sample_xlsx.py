#!/usr/bin/env python3
"""生成一个 Excel 示例文件，用来演示 /api/v1/import 的表格导入。

    python3 scripts/make_sample_xlsx.py [输出路径]

默认输出到 samples/sample_inventory.xlsx。
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEADERS = ["名称", "数量", "位置", "分类", "规格", "别名", "备注"]

ROWS = [
    ["STM32F103C8T6", 25, "A柜-1层-盒3", "STM元器件", "LQFP48", "F103C8, STM32F103", "蓝药丸板用"],
    ["0.1uF 50V MLCC", 500, "A柜-2层-盒1", "STM元器件", "0805", "104, 100nF", "退耦电容"],
    ["10kΩ 1% 电阻", 1000, "A柜-2层-盒2", "STM元器件", "0603", "10k, 103", ""],
    ["NE555", 30, "C柜-1层", "STM元器件", "DIP-8", "555, 定时器", ""],
    ["AMS1117-3.3", 50, "C柜-1层", "STM元器件", "SOT-223", "1117", "LDO"],
    ["1N4148", 200, "C柜-2层", "STM元器件", "DO-35", "4148", "开关二极管"],
    ["LM358 双运放", 12, "C柜-2层", "STM元器件", "SOIC-8", "LM358", ""],
    ["16MHz 晶振", 40, "C柜-3层", "STM元器件", "49S", "16M", ""],
    ["Steam 50元充值卡", 10, "B柜-抽屉1", "Steam游戏卡", "50元", "50元卡", ""],
    ["Steam 100元充值卡", 5, "B柜-抽屉2", "Steam游戏卡", "100元", "100元卡", ""],
]


def build(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "库存清单"

    sheet.append(HEADERS)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row in ROWS:
        sheet.append(row)

    for index, width in enumerate([28, 10, 18, 12, 14, 26, 20], start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"

    # 第二个工作表故意用不同的表头措辞，验证表头别名识别
    sheet2 = workbook.create_sheet("游戏卡")
    sheet2.append(["品名", "库存", "存放位置", "封装", "标签"])
    sheet2.append(["Steam 200元充值卡", 3, "B柜-抽屉3", "200元", "200元卡"])
    for index, width in enumerate([28, 10, 18, 12, 20], start=1):
        sheet2.column_dimensions[chr(64 + index)].width = width

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def main() -> int:
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "samples" / "sample_inventory.xlsx"
    saved = build(target)
    print(f"已生成：{saved}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
