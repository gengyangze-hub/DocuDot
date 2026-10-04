"""导入：规范化 Markdown / TXT / CSV / Excel → 库存记录。

容错是这里的第一原则。支持四种写法（可混用）：

1. **Markdown 表格**（推荐，也是 AI 预处理的输出格式）

   .. code-block:: markdown

      ## STM元器件

      | 名称 | 数量 | 位置 | 规格 | 别名 | 备注 |
      | --- | --- | --- | --- | --- | --- |
      | STM32F103C8T6 | 25 | A柜-1层-盒3 | LQFP48 | F103C8, STM32F103 | 蓝药丸用 |

2. **行内式条目**：``- STM32F103C8T6 × 25 @ A柜-1层 #F103C8 (LQFP48)``
3. **键值块**：``### 名称`` + 若干 ``- 数量: 25`` / ``位置: A柜``
4. **分隔符文本 / Excel**：表头行自动识别中英文别名

表头缺失时按固定列序解析：``名称 数量 位置 分类 规格 别名 备注 单位``。
"""

from __future__ import annotations

import csv
import io
import json
import logging
import re
from typing import Any, Iterable, Sequence

from ..config import Settings
from ..errors import ValidationFailed
from ..models import ImportPreviewOut, ImportResultOut, ImportRow, ItemOut
from ..repository import Repository
from ..utils import fmt_qty, now_iso
from .categorize import (
    DEFAULT_CATEGORY,
    category_from_heading,
    category_label,
    guess_category,
    normalize_category_code,
)

logger = logging.getLogger(__name__)

# --------------------------------------------------------------------------- #
# 表头别名
# --------------------------------------------------------------------------- #

HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "name": ("名称", "品名", "物品", "物品名称", "元件", "元件名", "型号", "器件", "name", "item", "part"),
    "quantity": ("数量", "数目", "库存", "库存数", "库存量", "个数", "件数", "qty", "quantity", "count", "stock"),
    "location": ("位置", "存储位置", "存放位置", "库位", "货架", "货架位", "柜子", "存放地", "location", "bin", "shelf"),
    "category": ("分类", "类别", "类型", "种类", "category", "type", "kind"),
    "spec": ("规格", "封装", "型号规格", "参数", "spec", "package", "footprint"),
    "aliases": ("别名", "标签", "别称", "又叫", "alias", "aliases", "tags", "tag"),
    "note": ("备注", "说明", "注释", "note", "remark", "comment"),
    "unit": ("单位", "计量单位", "unit"),
}

POSITIONAL_ORDER = ("name", "quantity", "location", "category", "spec", "aliases", "note", "unit")

_NUMBER_RE = re.compile(r"-?\d+(?:\.\d+)?")
_TABLE_SEP_RE = re.compile(r"^\s*\|?[\s:|-]+\|[\s:|-]*$")
_HEADING_RE = re.compile(r"^(#{1,6})\s*(.+?)\s*$")
_BULLET_RE = re.compile(r"^\s*(?:[-*+•]|\d+[.)])\s+(.*\S)\s*$")
_KV_RE = re.compile(r"^\s*(?:[-*+•]\s*)?([^:：]{1,12})\s*[:：]\s*(.*\S)\s*$")
#: 水平分隔线（``---`` / ``***`` / ``===``）
_HR_RE = re.compile(r"[-*_=]{3,}")
#: 行内条目必须带至少一个结构化标记，否则视为说明性文字。
#: 这条守卫用于挡掉「> 使用说明…」这类散落在正文里的句子被误当成物品。
_INLINE_ENTRY_MARKER = re.compile(
    r"[×✕@#（(]|\d\s*[x*×]\s*\d|\d+\s*(?:个|件|张|片|只|根|套|台|块|包|枚)"
)

_QUANTITY_HINT = re.compile(
    # ``×`` 前不能紧跟字母数字：否则尺寸里的 ``5x11mm`` 会被当成「乘 11」。
    # 实测踩到的坑：``10uF 25V（铝电解 5x11mm） × 200`` 解析出数量 11。
    r"(?:(?<![0-9A-Za-z])[×x✕*]\s*(\d+(?:\.\d+)?))"
    r"|(?:(\d+(?:\.\d+)?)\s*(?:个|件|张|片|只|根|套|台|块|包|枚))",
    re.IGNORECASE,
)
#: 行首的列表序号（``1. `` ``2)`` ``3、`` ``(4)`` ``①`` ``- ``）——是排版，不是名称的一部分。
#: 点号后**必须**有空白，否则 ``4.7kΩ`` 会被截成 ``7kΩ``；
#: 顿号/括号这类后面本来就不写空格，所以不要求。
_LIST_MARKER = re.compile(
    r"^\s*(?:"
    r"\d{1,3}\s*[、)）:：]\s*"
    r"|\d{1,3}\s*\.\s+"
    r"|[(（]\s*\d{1,3}\s*[)）]\s*"
    r"|[①-⑳]\s*"
    r"|[-*•·]\s+"
    r")"
)
#: 裸数量：``名称 50 @位置`` 这种没有 × 也没有单位的写法。
#: 三个守卫缺一不可：
#:   * 数字前必须是空白或行首 —— 否则 ``#555`` 这种别名里的数字会被吃掉
#:   * 数字后必须紧跟 ``@ # ( （`` 之一
#:   * 不能有前导零 —— ``0805`` 是封装，不是数量
_BARE_QUANTITY = re.compile(r"(?:^|\s)([1-9]\d*(?:\.\d+)?)\s*(?=[@#（(])")
#: 行尾裸数量：``AMS1117-3.3 5`` —— 最后一个独立数字就是数量。
#: 不能只在 ``_stock`` 里补，否则「名称里带小数点」（AMS1117-3.3）会被截错。
_BARE_QUANTITY_TAIL = re.compile(r"(?:^|\s)([1-9]\d*(?:\.\d+)?)\s*$")
_INLINE_LOCATION = re.compile(r"@\s*([^#@()（）]+?)\s*(?=#|\(|（|$)")
_INLINE_ALIAS = re.compile(r"#([^\s#@()（）]+)")
_INLINE_SPEC = re.compile(r"[（(]([^()（）]*)[）)]")
_INLINE_KEY_LOCATION = re.compile(r"(?:位置|库位|存放)\s*[:：]?\s*([^\s#@()（）]+)")
#: 名称里至少要有字母/数字/汉字，纯标点（比如单独一个 ``#``）不算条目
_NAME_HAS_CONTENT = re.compile(r"[0-9A-Za-z\u4e00-\u9fff]")
#: 纯数字单元格
_NUMERIC_CELL = re.compile(r"^-?\d+(?:\.\d+)?$")


# --------------------------------------------------------------------------- #
# 小工具
# --------------------------------------------------------------------------- #


def _clean_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _split_aliases(raw: str) -> list[str]:
    if not raw:
        return []
    parts = re.split(r"[,，、/|;；]+", raw)
    seen: set[str] = set()
    result: list[str] = []
    for part in parts:
        text = part.strip().strip("#").strip()
        if text and text not in seen:
            seen.add(text)
            result.append(text)
    return result


def _parse_quantity(raw: str, issues: list[str], row_number: int) -> float:
    text = _clean_cell(raw)
    if not text:
        issues.append(f"第 {row_number} 行没有数量，按 0 处理")
        return 0.0
    match = _NUMBER_RE.search(text)
    if not match:
        issues.append(f"第 {row_number} 行数量「{text}」无法识别，按 0 处理")
        return 0.0
    value = float(match.group())
    if value < 0:
        issues.append(f"第 {row_number} 行数量为负数（{value:g}），已取绝对值")
        value = abs(value)
    return value


def _normalize_category(raw: str) -> str | None:
    """分类单元格 → 分类名。

    **没有内置分类表了** —— 用户在表里写什么分类就是什么分类；
    空值或写不出有效分类名时返回 ``None``，让调用方回退到标题 / 默认值。
    """
    text = (raw or "").strip()
    if not text:
        return None
    candidate = normalize_category_code(text)
    return None if candidate == DEFAULT_CATEGORY else candidate


def _map_header(cells: Sequence[str]) -> tuple[dict[int, str], list[str]]:
    """把表头单元格映射成字段名。返回 ``(列索引→字段, 识别到的表头文本)``。"""
    mapping: dict[int, str] = {}
    detected: list[str] = []
    for index, cell in enumerate(cells):
        text = _clean_cell(cell).casefold()
        if not text:
            continue
        for field, aliases in HEADER_ALIASES.items():
            if field in mapping.values():
                continue
            if any(text == alias.casefold() or alias.casefold() in text for alias in aliases):
                mapping[index] = field
                detected.append(_clean_cell(cell))
                break
    return mapping, detected


def decide_header(cells: Sequence[str]) -> tuple[dict[int, str], list[str], bool]:
    """判断某一行是不是表头，返回 ``(列映射, 命中的表头文本, 是否是表头)``。

    两个条件同时成立才算表头：

    1. 认出 **≥2 个字段**（一个字段的偶合不算）
    2. **没有任何单元格是纯数字** —— 表头不会叫「25」，而数据行经常是

    第 2 条是必需的：表头识别用的是「包含」匹配，而中文里
    「器件」是「元器件」的子串，``| STM32F103C8T6 | 25 | A柜 | STM元器件 | … |``
    会被误判成表头，把整行数据吞掉。
    """
    mapping, detected = _map_header(cells)
    if len(mapping) < 2:
        return mapping, detected, False
    if any(_NUMERIC_CELL.match(_clean_cell(cell)) for cell in cells):
        return mapping, detected, False
    return mapping, detected, True


def _row_from_cells(
    cells: Sequence[str],
    mapping: dict[int, str],
    *,
    row_number: int,
    default_category: str | None,
    heading_category: str | None,
) -> ImportRow | None:
    data: dict[str, str] = {}
    if mapping:
        for index, field in mapping.items():
            if index < len(cells):
                data[field] = _clean_cell(cells[index])
    else:
        for index, field in enumerate(POSITIONAL_ORDER):
            if index < len(cells):
                data[field] = _clean_cell(cells[index])

    name = data.get("name", "").strip()
    if not name:
        return None

    issues: list[str] = []
    quantity = _parse_quantity(data.get("quantity", ""), issues, row_number)
    category = _normalize_category(data.get("category", "")) or heading_category or guess_category(name)
    row = ImportRow(
        name=name,
        quantity=quantity,
        location=data.get("location", "").strip(),
        category=category,  # type: ignore[arg-type]
        spec=data.get("spec", "").strip(),
        unit=data.get("unit", "").strip(),
        aliases=_split_aliases(data.get("aliases", "")),
        note=data.get("note", "").strip(),
        row_number=row_number,
        issues=issues,
    )
    if default_category and heading_category is None and not _normalize_category(data.get("category", "")):
        row.category = default_category  # type: ignore[assignment]
    return row


# --------------------------------------------------------------------------- #
# Markdown / 文本
# --------------------------------------------------------------------------- #


def _iter_sections(text: str) -> Iterable[tuple[int, str, list[str]]]:
    """按 Markdown 标题切分，返回 ``(标题层级, 标题, 正文行)``。

    层级很关键：``##`` 是分类章节，``###`` 是条目名。
    """
    level = 0
    heading = ""
    buffer: list[str] = []
    for line in text.splitlines():
        match = _HEADING_RE.match(line)
        # ``#`` 也可能是行内条目的别名标记，仅当整行只有标题时才算标题
        if match and not _BULLET_RE.match(line):
            if buffer or heading:
                yield level, heading, buffer
                buffer = []
            level = len(match.group(1))
            heading = match.group(2).strip()
            continue
        buffer.append(line)
    if buffer or heading:
        yield level, heading, buffer


def _parse_table_block(lines: Sequence[str], row_start: int) -> tuple[list[list[str]], int]:
    """把连续的表格行拆成单元格矩阵。"""
    matrix: list[list[str]] = []
    index = 0
    while index < len(lines):
        line = lines[index]
        if "|" not in line:
            break
        if _TABLE_SEP_RE.match(line) and matrix:
            index += 1
            continue
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        matrix.append(cells)
        index += 1
    return matrix, row_start + index


#: 可直接**从名称里提出来当规格**的封装码（比 ``_PACKAGE_TAG`` 严：
#: 不收裸数字，免得把 ``USB线 2.0`` 的 ``2.0`` 提成规格）
_PROMOTABLE_PACKAGE = re.compile(
    r"^(?:0[46]03|0805|1206|1210|2010|2512|直插|贴片|插件|smd|mlcc|瓷片|瓷介|铝电解|电解|钽|"
    r"\d+(?:\.\d+)?mm|\d+x\d+(?:\.\d+)?mm|"
    r"(?:dip|sop|soic|sot|qfn|lqfp|dfn|bga|to|do|sod)-?\d*)$",
    re.IGNORECASE,
)


def _promote_trailing_package(name: str) -> tuple[str, str]:
    """把名称末尾裸写的封装码提出来当规格：``100Ω电容 0805`` → ``("100Ω电容", "0805")``。

    不提的话，机器人会再追问一次封装，用户答完 ``0805`` 后名称与规格就重复了。
    只在「还剩别的词」时才提，避免把单独一行的 ``0805`` 拆成空名称。
    """
    tokens = name.split()
    if len(tokens) < 2:
        return name, ""
    last = tokens[-1]
    if not _PROMOTABLE_PACKAGE.match(last):
        return name, ""
    head = " ".join(tokens[:-1]).strip()
    if not head or not _NAME_HAS_CONTENT.search(head):
        return name, ""
    return head, last


def parse_inline_entry(text: str, row_number: int = 0, heading_category: str | None = None) -> ImportRow | None:
    """``名称 × 25 @ 位置 #别名 (规格)`` 形式的行内条目。

    QQ 命令（``入库 STM32F103C8T6 25 @A柜``）也复用这套解析。
    """
    original = _LIST_MARKER.sub("", text.strip())
    if not original:
        return None

    aliases: list[str] = []
    for match in _INLINE_ALIAS.finditer(original):
        aliases.extend(_split_aliases(match.group(1)))
    spec = ""
    spec_match = _INLINE_SPEC.search(original)
    if spec_match:
        spec = spec_match.group(1).strip()

    location = ""
    location_match = _INLINE_LOCATION.search(original) or _INLINE_KEY_LOCATION.search(original)
    if location_match:
        location = location_match.group(1).strip()

    quantity = 0.0
    quantity_match = _QUANTITY_HINT.search(original)
    if quantity_match:
        quantity = float(quantity_match.group(1) or quantity_match.group(2))
    else:
        bare = _BARE_QUANTITY.search(original) or _BARE_QUANTITY_TAIL.search(original)
        if bare:
            quantity = float(bare.group(1))

    # 名称 = 原串去掉所有已识别的片段。
    # 顺序很关键：先剔除 ``× 25`` 这种带算符的数量（否则会留下一个孤零零的 ×），
    # 再剔除裸数量（它依赖后面还跟着 @/#/括号 或位于行尾才能被识别），
    # 最后才清别名、规格与位置。
    name = original
    for pattern in (
        _QUANTITY_HINT,
        _BARE_QUANTITY,
        _BARE_QUANTITY_TAIL,
        _INLINE_ALIAS,
        _INLINE_SPEC,
        _INLINE_LOCATION,
        _INLINE_KEY_LOCATION,
    ):
        name = pattern.sub(" ", name)
    name = re.sub(r"\s+", " ", name).strip(" -#\t@")
    if not name or len(name) > 120:
        return None
    if not _NAME_HAS_CONTENT.search(name):
        # 纯标点（比如单独一行 ``#``）不是物品
        return None

    # 裸写的封装码提到 spec（``100Ω电容 0805 100`` → 名称 ``100Ω电容``、规格 ``0805``）
    if not spec:
        name, promoted = _promote_trailing_package(name)
        spec = promoted

    issues: list[str] = []
    if quantity <= 0:
        issues.append(f"第 {row_number} 行未识别到数量，按 0 处理")

    category = heading_category or guess_category(name)
    return ImportRow(
        name=name,
        quantity=quantity,
        location=location,
        category=category,  # type: ignore[arg-type]
        spec=spec,
        aliases=aliases,
        row_number=row_number,
        issues=issues,
    )


def parse_markdown(text: str, *, default_category: str | None = None) -> ImportPreviewOut:
    """解析 Markdown / 纯文本内容。"""
    rows: list[ImportRow] = []
    warnings: list[str] = []
    detected_columns: list[str] = []
    line_no = 0

    for level, heading, body in _iter_sections(text):
        # ``###`` 及更深是**条目名**，本身不是分类（下面那段旧代码也这么约定）。
        # 只有 ``##`` 才是分类章节 —— 否则「### Steam 50元充值卡」会被当成一个分类。
        heading_category = (
            category_from_heading(heading) if heading and level <= 2 else None
        )
        index = 0
        pending_block: list[str] = []
        pending_name: str | None = None
        # ``###`` 及更深的标题就是条目名（``##`` 才是分类章节）
        if level >= 3 and heading:
            pending_name = heading
            pending_block = [f"名称: {heading}"]

        def flush_block() -> None:
            nonlocal pending_block, pending_name
            if not pending_block:
                pending_name = None
                return
            data: dict[str, str] = {}
            for candidate in pending_block:
                kv = _KV_RE.match(candidate)
                if not kv:
                    continue
                key, value = kv.group(1).strip().casefold(), kv.group(2).strip()
                for field, aliases in HEADER_ALIASES.items():
                    if any(key == alias.casefold() or alias.casefold() in key for alias in aliases):
                        data[field] = value
                        break
            if pending_name and "name" not in data:
                data["name"] = pending_name
            # 除了名字之外至少还要有一个真实字段，否则视为说明文字。
            # 否则「### 某个小标题 + 一段话」会凭空变成一条库存记录。
            has_payload = any(field != "name" for field in data)
            if data.get("name") and has_payload:
                issues: list[str] = []
                quantity = _parse_quantity(data.get("quantity", ""), issues, line_no)
                category = _normalize_category(data.get("category", "")) or heading_category or guess_category(
                    data["name"]
                )
                rows.append(
                    ImportRow(
                        name=data["name"],
                        quantity=quantity,
                        location=data.get("location", ""),
                        category=category,  # type: ignore[arg-type]
                        spec=data.get("spec", ""),
                        unit=data.get("unit", ""),
                        aliases=_split_aliases(data.get("aliases", "")),
                        note=data.get("note", ""),
                        row_number=line_no,
                        issues=issues,
                    )
                )
            pending_block = []
            pending_name = None

        while index < len(body):
            line = body[index]
            line_no += 1
            stripped = line.strip()

            if not stripped:
                flush_block()
                index += 1
                continue

            # ---- 引用块 / 分隔线：说明性内容，直接跳过 ----
            if stripped.startswith(">") or _HR_RE.fullmatch(stripped):
                flush_block()
                index += 1
                continue

            # ---- 表格 ----
            if "|" in stripped and stripped.count("|") >= 2:
                flush_block()
                matrix, consumed = _parse_table_block(body[index:], line_no)
                index += consumed
                line_no += consumed
                if matrix:
                    header_map, detected, has_header = decide_header(matrix[0])
                    detected_columns.extend(detected if has_header else [])
                    data_rows = matrix[1:] if has_header else matrix
                    if not has_header:
                        warnings.append(
                            f"「{heading or '未命名章节'}」的表格没有识别到表头，"
                            f"按固定列序（{'/'.join(POSITIONAL_ORDER)}）解析"
                        )
                    for offset, cells in enumerate(data_rows, start=1):
                        row = _row_from_cells(
                            cells,
                            header_map if has_header else {},
                            row_number=line_no - len(data_rows) + offset,
                            default_category=default_category,
                            heading_category=heading_category,
                        )
                        if row:
                            rows.append(row)
                continue

            # ---- 键值行 / 行内条目 ----
            bullet = _BULLET_RE.match(stripped)
            content = bullet.group(1) if bullet else stripped
            if _KV_RE.match(content) and not _INLINE_ALIAS.search(content):
                pending_block.append(content)
                index += 1
                continue

            # 非列表项且没有任何结构化标记 → 当说明文字跳过（不产生条目）
            if bullet is None and not _INLINE_ENTRY_MARKER.search(content):
                flush_block()
                index += 1
                continue

            flush_block()
            row = parse_inline_entry(content, line_no, heading_category)
            if row:
                if default_category and heading_category is None and row.category == "other":
                    row.category = default_category  # type: ignore[assignment]
                rows.append(row)
            index += 1

        flush_block()

    for row in rows:
        warnings.extend(row.issues)

    return ImportPreviewOut(
        filename="",
        fmt="markdown",
        rows=rows,
        total=len(rows),
        warnings=warnings,
        detected_columns=sorted(set(detected_columns)),
    )


def parse_delimited(text: str, *, delimiter: str | None = None, default_category: str | None = None) -> ImportPreviewOut:
    """解析 CSV / TSV 文本。"""
    sample = text[:4096]
    if delimiter is None:
        delimiter = "\t" if sample.count("\t") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    matrix = [[cell.strip() for cell in row] for row in reader if any(cell.strip() for cell in row)]
    if not matrix:
        return ImportPreviewOut(filename="", fmt="delimited", rows=[], total=0)

    header_map, detected, has_header = decide_header(matrix[0])
    data_rows = matrix[1:] if has_header else matrix
    warnings: list[str] = []
    if not has_header:
        warnings.append("未识别到表头，按固定列序解析")
    rows: list[ImportRow] = []
    for offset, cells in enumerate(data_rows, start=1):
        row = _row_from_cells(
            cells,
            header_map,
            row_number=offset + 1,
            default_category=default_category,
            heading_category=None,
        )
        if row:
            rows.append(row)
    warnings.extend(issue for row in rows for issue in row.issues)
    return ImportPreviewOut(
        filename="",
        fmt="delimited",
        rows=rows,
        total=len(rows),
        warnings=warnings,
        detected_columns=sorted(set(detected)),
    )


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #


def parse_excel(data: bytes, *, default_category: str | None = None) -> ImportPreviewOut:
    """解析 ``.xlsx`` / ``.xlsm``（需要 openpyxl）。"""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ValidationFailed("解析 Excel 需要 openpyxl，请先 pip install openpyxl") from exc

    try:
        workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl 会抛各种底层异常
        # 上传一个不是 xlsx 的文件（改名、损坏、zip 炸弹）不能变成 500。
        # 实测：传 "not really an xlsx" 会抛 zipfile.BadZipFile 直接冒到接口层。
        raise ValidationFailed(
            "这不是有效的 .xlsx 文件（可能已损坏，或只是把别的文件改了后缀）。"
            "Excel 请另存为 .xlsx；旧版 .xls 请先另存。"
        ) from exc
    try:
        return _parse_workbook(workbook, default_category=default_category)
    finally:
        try:
            workbook.close()
        except Exception:  # pragma: no cover - 关不掉也不该影响结果
            logger.debug("关闭 Excel 工作簿时出错", exc_info=True)


def _parse_workbook(workbook: Any, *, default_category: str | None = None) -> ImportPreviewOut:
    rows: list[ImportRow] = []
    warnings: list[str] = []
    detected: list[str] = []

    for sheet in workbook.worksheets:
        sheet_category = category_from_heading(sheet.title)
        matrix: list[list[str]] = []
        for raw_row in sheet.iter_rows(values_only=True):
            cells = [_clean_cell(cell) for cell in raw_row]
            if any(cells):
                matrix.append(cells)
        if not matrix:
            continue

        # 前 10 行里找表头：命中至少 2 个字段才算
        header_index = -1
        header_map: dict[int, str] = {}
        for candidate in range(min(10, len(matrix))):
            mapping, found, ok = decide_header(matrix[candidate])
            if ok:
                header_index = candidate
                header_map = mapping
                detected.extend(found)
                break

        if header_index < 0:
            warnings.append(f"工作表「{sheet.title}」没识别到表头，整表按固定列序解析")
            data_rows = matrix
            offset = 0
        else:
            data_rows = matrix[header_index + 1 :]
            offset = header_index + 1

        for position, cells in enumerate(data_rows, start=offset + 1):
            row = _row_from_cells(
                cells,
                header_map,
                row_number=position,
                default_category=default_category,
                heading_category=sheet_category,
            )
            if row:
                row.issues = [f"[{sheet.title}] {issue}" for issue in row.issues]
                rows.append(row)

    warnings.extend(issue for row in rows for issue in row.issues)
    return ImportPreviewOut(
        filename="",
        fmt="excel",
        rows=rows,
        total=len(rows),
        warnings=warnings,
        detected_columns=sorted(set(detected)),
    )


# --------------------------------------------------------------------------- #
# JSON
# --------------------------------------------------------------------------- #

#: component-inventory 的粗粒度库存等级 → 代表性数量。
#: 它是「少 / 中 / 多」的等级而不是精确件数，所以**原始等级会写进备注**，
#: 免得映射出来的数字被当成真实库存。
COARSE_STOCK_LEVELS: dict[int, float] = {1: 10.0, 2: 50.0, 3: 200.0}
_COARSE_LABELS: dict[int, str] = {1: "少", 2: "中", 3: "多"}

#: 看起来像「封装 / 规格」的标签（0805、DIP-8、MLCC、瓷片、2.54、5x11mm…）
_PACKAGE_TAG = re.compile(
    r"^(?:0[46]03|0805|1206|1210|2010|2512|直插|贴片|插件|smd|mlcc|瓷片|瓷介|铝电解|电解|钽|"
    r"\d+(?:\.\d+)?(?:mm)?|\d+x\d+(?:\.\d+)?mm|"
    r"(?:dip|sop|soic|sot|qfn|lqfp|dfn|bga|to|do|sod)-?\d*)$",
    re.IGNORECASE,
)

#: JSON 里可能包裹记录数组的键
_JSON_LIST_KEYS = ("items", "data", "records", "rows", "list", "entries", "inventory", "components")


def _json_scalar(value: Any) -> str:
    """把 JSON 的标量/数组值转成字符串（数组用逗号连接）。"""
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, (list, tuple)):
        return ", ".join(part for part in (_json_scalar(item) for item in value) if part)
    if isinstance(value, dict):
        return ""
    return str(value).strip()


def _json_field(key: str) -> str | None:
    """把 JSON 的键映射成字段名（复用表头别名表）。"""
    text = _clean_cell(key).casefold()
    if not text:
        return None
    for field, aliases in HEADER_ALIASES.items():
        if any(text == alias.casefold() or alias.casefold() in text for alias in aliases):
            return field
    return None


def _json_records(payload: Any) -> list[Any]:
    """从任意 JSON 结构里找出记录数组。"""
    if isinstance(payload, list):
        return payload
    if not isinstance(payload, dict):
        return []
    for key in _JSON_LIST_KEYS:
        value = payload.get(key)
        if isinstance(value, list):
            return value
    lists = [value for value in payload.values() if isinstance(value, list)]
    return lists[0] if len(lists) == 1 else []


def _component_inventory_row(record: dict, tags: list[str], row_number: int) -> ImportRow:
    """把 component-inventory 的一条记录映射成本系统的行。

    它的「身份」是 ``tags``（如 ``C / 1uF / 16V / MLCC / 0805``），
    库存是 ``stock.mode = coarse`` 的等级（少/中/多），没有位置概念。
    """
    issues: list[str] = []
    note_bits: list[str] = []
    quantity = 0.0

    stock = record.get("stock")
    if isinstance(stock, dict):
        exact = stock.get("quantity", stock.get("qty"))
        level_raw = stock.get("level")
        if exact is None and level_raw is not None:
            try:
                level = int(level_raw)
            except (TypeError, ValueError):
                level = 0
            quantity = COARSE_STOCK_LEVELS.get(level, 0.0)
            note_bits.append(f"原粗粒度库存 level {level}（{_COARSE_LABELS.get(level, '未知')}）")
            if level not in COARSE_STOCK_LEVELS:
                issues.append(f"未知的粗粒度等级 {level_raw}，数量按 0 处理")
        elif exact is not None:
            try:
                quantity = float(exact)
            except (TypeError, ValueError):
                issues.append(f"数量 {exact!r} 无法识别，按 0 处理")

    raw_note = record.get("note")
    if raw_note:
        note_bits.append(str(raw_note).strip())

    packages = [tag for tag in tags if _PACKAGE_TAG.match(tag)]
    # 封装单独进 spec，名称里不重复写（全是封装时退回原样，避免空名称）
    base = [tag for tag in tags if tag not in packages]
    return ImportRow(
        name=" ".join(base) or " ".join(tags),
        quantity=quantity,
        location=_json_scalar(record.get("location")),
        category=DEFAULT_CATEGORY,  # 这个导出格式自己不带分类，交给 AI 归类
        spec=" ".join(packages),
        aliases=[],
        note="；".join(bit for bit in note_bits if bit),
        row_number=row_number,
        issues=issues,
    )


def parse_json(text: str, default_category: str | None = None) -> ImportPreviewOut:
    """解析 JSON 清单。

    兼容三类结构：

    * 顶层数组：``[{"名称": "NE555", "数量": 30}]``
    * 包一层：``{"items": [...]}``（``components`` / ``records`` / ``data`` 等同样认）
    * **component-inventory 导出**：``{"components": [{"tags": [...], "stock": {...}}]}``
    """
    try:
        payload = json.loads(text)
    except json.JSONDecodeError as exc:
        raise ValidationFailed(f"JSON 解析失败：{exc.msg}（第 {exc.lineno} 行）") from exc

    records = _json_records(payload)
    if not records:
        raise ValidationFailed(
            "JSON 里没找到记录数组（支持顶层数组，或 items / components / records / data 字段）"
        )

    rows: list[ImportRow] = []
    warnings: list[str] = []
    coarse_seen = False

    for index, record in enumerate(records, start=1):
        if isinstance(record, str):
            name = record.strip()
            if name:
                rows.append(
                    ImportRow(
                        name=name,
                        quantity=0,
                        category=default_category or guess_category(name),
                        row_number=index,
                    )
                )
            continue
        if not isinstance(record, dict):
            warnings.append(f"第 {index} 条不是对象，已跳过")
            continue

        tags = record.get("tags")
        if isinstance(tags, list) and tags:
            clean = [str(tag).strip() for tag in tags if str(tag).strip()]
            if clean:
                row = _component_inventory_row(record, clean, index)
                coarse_seen = coarse_seen or "原粗粒度库存" in row.note
                rows.append(row)
                continue

        data: dict[str, str] = {}
        for key, value in record.items():
            field = _json_field(str(key))
            if field and field not in data:
                data[field] = _json_scalar(value)

        name = data.get("name", "").strip()
        if not name:
            warnings.append(f"第 {index} 条没有名称，已跳过")
            continue

        issues: list[str] = []
        quantity = _parse_quantity(data.get("quantity", ""), issues, index)
        category = (
            _normalize_category(data.get("category", "")) or default_category or guess_category(name)
        )
        rows.append(
            ImportRow(
                name=name,
                quantity=quantity,
                location=data.get("location", "").strip(),
                category=category,  # type: ignore[arg-type]
                spec=data.get("spec", "").strip(),
                unit=data.get("unit", "").strip(),
                aliases=_split_aliases(data.get("aliases", "")),
                note=data.get("note", "").strip(),
                row_number=index,
                issues=issues,
            )
        )

    if not rows:
        raise ValidationFailed("JSON 里没有可导入的条目")

    if coarse_seen:
        warnings.insert(
            0,
            "检测到 component-inventory 的粗粒度库存，已按 1→10 / 2→50 / 3→200 映射，"
            "原始等级写在备注里",
        )

    return ImportPreviewOut(filename="", fmt="json", rows=rows, total=len(rows), warnings=warnings)


# --------------------------------------------------------------------------- #
# 入口
# --------------------------------------------------------------------------- #


def detect_format(filename: str, data: bytes | None = None, text: str | None = None) -> str:
    lowered = (filename or "").casefold()
    if lowered.endswith((".xlsx", ".xlsm")):
        return "excel"
    if lowered.endswith(".xls"):
        raise ValidationFailed("不支持旧版 .xls，请在 Excel 里另存为 .xlsx 或 CSV")
    if lowered.endswith(".csv"):
        return "csv"
    if lowered.endswith((".md", ".markdown")):
        return "markdown"
    if lowered.endswith(".json"):
        return "json"
    if lowered.endswith((".txt", ".text", ".log")):
        return "text"
    if text:
        # 没有扩展名时按内容嗅探
        if text.lstrip().startswith(("{", "[")):
            return "json"
        if re.search(r"^\s*\|.*\|\s*$", text, re.MULTILINE):
            return "markdown"
        return "text"
    return "text"


def parse_input(
    filename: str,
    *,
    data: bytes | None = None,
    text: str | None = None,
    default_category: str | None = None,
) -> ImportPreviewOut:
    """统一入口：自动判格式并解析。"""
    fmt = detect_format(filename, data, text)
    if fmt == "excel":
        if data is None:
            raise ValidationFailed("Excel 导入需要上传文件二进制内容")
        preview = parse_excel(data, default_category=default_category)
    else:
        if text is None:
            if data is None:
                raise ValidationFailed("没有可解析的内容")
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                try:
                    text = data.decode("gbk")
                except UnicodeDecodeError as exc:
                    raise ValidationFailed("文件编码无法识别，请另存为 UTF-8") from exc
        if fmt == "json":
            preview = parse_json(text, default_category=default_category)
        elif fmt == "csv":
            preview = parse_delimited(text, delimiter=",", default_category=default_category)
        elif fmt == "markdown":
            preview = parse_markdown(text, default_category=default_category)
        else:
            # 文本里也可能有表格或键值块
            preview = parse_delimited(text, default_category=default_category) if "\t" in text[:2000] else parse_markdown(
                text, default_category=default_category
            )
    preview.filename = filename
    preview.fmt = fmt
    return preview


# --------------------------------------------------------------------------- #
# 落库
# --------------------------------------------------------------------------- #


class ImportService:
    """把解析结果写进数据库，并记录导入批次与流水。"""

    def __init__(self, repo: Repository, settings: Settings) -> None:
        self.repo = repo
        self.settings = settings

    def commit(
        self,
        preview: ImportPreviewOut,
        *,
        mode: str = "merge",
        operator: str = "",
        source: str = "import",
        dry_run: bool = False,
    ) -> ImportResultOut:
        rows = [row for row in preview.rows if row.name.strip()]
        if not rows:
            raise ValidationFailed("没有解析到任何有效条目")

        created = updated = skipped = 0
        items: list[ItemOut] = []
        warnings = list(preview.warnings)

        if mode == "replace" and not dry_run:
            categories = {row.category.value if hasattr(row.category, "value") else row.category for row in rows}
            for category in categories:
                removed = self.repo.delete_by_category(category)
                if removed:
                    warnings.append(f"replace 模式：已清空「{category}」下 {removed} 条旧记录")

        for row in rows:
            category = row.category.value if hasattr(row.category, "value") else str(row.category)
            existing = self.repo.get_item_by_key(row.name, row.location, row.spec)
            if existing and mode == "add":
                before = existing.quantity
                after = before + row.quantity
                if not dry_run:
                    self.repo.set_quantity(existing.id, after, operator=operator)
                    self.repo.record_movement(
                        item_id=existing.id,
                        item_name=existing.name,
                        action="import",
                        delta=row.quantity,
                        quantity_before=before,
                        quantity_after=after,
                        location=existing.location,
                        operator=operator,
                        source=source,
                        note=f"导入并入（{preview.filename}）",
                    )
                updated += 1
                record = self.repo.get_item(existing.id)
                if record:
                    items.append(record.to_out())
                continue

            if existing:
                if dry_run:
                    updated += 1
                    items.append(existing.to_out())
                    continue
                fields: dict[str, Any] = {
                    "category": category,
                    "quantity": row.quantity,
                    "unit": row.unit or existing.unit,
                    "note": row.note or existing.note,
                }
                if row.aliases:
                    merged = list(dict.fromkeys([*existing.aliases, *row.aliases]))
                    fields["aliases"] = merged
                before = existing.quantity
                self.repo.update_item(existing.id, fields=fields, operator=operator)
                self.repo.record_movement(
                    item_id=existing.id,
                    item_name=existing.name,
                    action="import",
                    delta=row.quantity - before,
                    quantity_before=before,
                    quantity_after=row.quantity,
                    location=existing.location,
                    operator=operator,
                    source=source,
                    note=f"导入覆盖（{preview.filename}）",
                )
                updated += 1
                record = self.repo.get_item(existing.id)
                if record:
                    items.append(record.to_out())
                continue

            if dry_run:
                created += 1
                continue
            record = self.repo.create_item(
                name=row.name,
                category=category,
                quantity=row.quantity,
                unit=row.unit,
                location=row.location,
                spec=row.spec,
                note=row.note,
                aliases=row.aliases,
                operator=operator,
            )
            self.repo.record_movement(
                item_id=record.id,
                item_name=record.name,
                action="import",
                delta=record.quantity,
                quantity_before=0,
                quantity_after=record.quantity,
                location=record.location,
                operator=operator,
                source=source,
                note=f"导入新建（{preview.filename}）",
            )
            created += 1
            items.append(record.to_out())

        status = "ok" if not warnings else "partial"
        batch_id = None
        if not dry_run:
            batch_id = self.repo.create_import_batch(
                filename=preview.filename or "inline",
                fmt=preview.fmt,
                mode=mode,
                status=status,
                total=len(rows),
                created_items=created,
                updated_items=updated,
                skipped=skipped,
                message="; ".join(warnings[:5]),
                operator=operator,
                source=source,
            )
            self.repo.audit(
                action="import.commit",
                actor=operator,
                actor_kind="qq" if source == "qq" else "api",
                target_type="import_batch",
                target_id=batch_id or "",
                detail={"filename": preview.filename, "total": len(rows), "created": created, "updated": updated},
            )

        return ImportResultOut(
            batch_id=batch_id,
            filename=preview.filename or "inline",
            fmt=preview.fmt,
            status="dry-run" if dry_run else status,
            total=len(rows),
            created=created,
            updated=updated,
            skipped=skipped,
            warnings=warnings,
            items=items[:50],
            message=(
                f"{'[试运行] ' if dry_run else ''}导入完成："
                f"共 {len(rows)} 条，新建 {created}，更新 {updated}"
                + (f"，{len(warnings)} 条提示" if warnings else "")
            ),
        )


def preview_to_markdown(preview: ImportPreviewOut, limit: int = 50) -> str:
    """把解析结果渲染成规范化 Markdown（便于人工复核 / 二次导入）。"""
    groups: dict[str, list[ImportRow]] = {}
    for row in preview.rows:
        category = row.category.value if hasattr(row.category, "value") else str(row.category)
        groups.setdefault(category, []).append(row)

    labels = {category: category_label(category) for category in groups}
    lines = ["# 库存清单", ""]
    for category, rows in groups.items():
        lines.append(f"## {labels.get(category, category)}")
        lines.append("")
        lines.append("| 名称 | 数量 | 位置 | 规格 | 别名 | 备注 |")
        lines.append("| --- | --- | --- | --- | --- | --- |")
        for row in rows[:limit]:
            aliases = ", ".join(row.aliases)
            lines.append(
                f"| {row.name} | {fmt_qty(row.quantity)} | {row.location} | {row.spec} | {aliases} | {row.note} |"
            )
        lines.append("")
    return "\n".join(lines)
